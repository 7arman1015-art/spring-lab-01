"""Create Lab 6 by patching a retained Lab 3 OOXML package.

The reference remains untouched. --prepare writes the distillation contract only.
--build is gated on the documented reference visual review and reuses the source
paragraph, table and drawing components without rebuilding opaque package parts.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from lxml import etree as ET
from PIL import Image

from evidence_metadata import english_screenshot

ROOT = Path(__file__).resolve().parents[2]
REPORT_DIR = ROOT / "reports" / "lab06"
REFERENCE_DIR = REPORT_DIR / "qa" / "reference"
REFERENCE = REFERENCE_DIR / "Lab03_Completed_Arman_Aruzhan.docx"
CONTRACT = REFERENCE_DIR / "artifact.md"
BASELINE = REFERENCE_DIR / "contract_baseline.json"
TEMPLATE = REPORT_DIR / "report_template.docx"
COMPONENTS = REPORT_DIR / "report_template_components.json"
PACKAGE_VERIFICATION = REPORT_DIR / "evidence" / "package_verification.json"
EVIDENCE = REPORT_DIR / "evidence"
OUTPUT = REPORT_DIR / "Laboratory_work_6_report.docx"
BUILD_LOG = ROOT / "target" / "memory-check" / "maven-package.log"
REPO_LINK = "https://github.com/7arman1015-art/spring-lab-01/tree/lab06"
REFERENCE_HASH = "1deb008df53cfec99d90f3a465f0b36256be21186f1f88c876c81147b3236d1d"
NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}
EDITABLE_PARTS = {"word/document.xml", "word/_rels/document.xml.rels",
                  "[Content_Types].xml", "docProps/core.xml"}
PROTOTYPES = {"institution": 0, "department": 1, "title": 3, "topic": 4,
              "metadata": 5, "objective_heading": 7, "body": 9,
              "task_heading": 11, "explanation": 24, "subheading": 28,
              "result": 35, "image": 40, "caption": 69,
              "compact_heading": 77, "compact_body": 78}


def tag(prefix: str, local: str) -> str:
    return "{" + NS[prefix] + "}" + local


def element(prefix: str, local: str, attrs: dict | None = None):
    out = ET.Element(tag(prefix, local))
    for name, value in (attrs or {}).items():
        out.set(tag(prefix, name), str(value))
    return out


def verified_reference() -> bytes:
    data = REFERENCE.read_bytes()
    if sha256(data).hexdigest() != REFERENCE_HASH:
        raise ValueError("Retained reference changed; repeat template distillation.")
    return data


def compact_xml(el) -> dict:
    if el is None:
        return {}
    return {ET.QName(child).localname:
            {ET.QName(key).localname: value for key, value in child.attrib.items()}
            for child in el}


def prepare_contract():
    verified_reference()
    doc = Document(REFERENCE)  # read-only inspection; never saved through python-docx
    with ZipFile(REFERENCE) as z:
        inventory = [{"part": name, "size": len(z.read(name)),
                      "sha256": sha256(z.read(name)).hexdigest(),
                      "classification": "editable" if name in EDITABLE_PARTS else "preserve-only"}
                     for name in z.namelist()]
        relations = ET.fromstring(z.read("word/_rels/document.xml.rels"))
        baseline = {
            "reference": str(REFERENCE), "sha256": REFERENCE_HASH,
            "part_inventory": inventory,
            "relationships": [dict(rel.attrib) for rel in relations],
            "section_xml": ET.tostring(doc.sections[0]._sectPr).decode(),
            "prototypes": {
                role: {"paragraph_index": i, "style_id": doc.paragraphs[i].style.style_id,
                       "pPr": compact_xml(doc.paragraphs[i]._p.pPr),
                       "paragraph_xml": ET.tostring(doc.paragraphs[i]._p).decode()}
                for role, i in PROTOTYPES.items()},
            "tables": [{"table_index": i, "rows": len(t.rows), "columns": len(t.columns),
                        "table_xml": ET.tostring(t._tbl).decode()} for i, t in enumerate(doc.tables)],
        }
    BASELINE.write_text(json.dumps(baseline, ensure_ascii=False, indent=2), encoding="utf-8")
    text = f"""# Lab 6 template execution contract

Reference visual gate: pending
Reference page count: pending canonical render and inspection
Reference render evidence: supplied by root after inspection of every page

## Reference authority

- Retained DOCX: `{REFERENCE}`
- SHA256: `{REFERENCE_HASH}`
- Read-only OOXML evidence: `{REFERENCE_DIR / 'inspection.json'}`
- Exact part, relationship, prototype and table inventory: `{BASELINE}`
- Builder: `{Path(__file__).resolve()}`
- Final output: `{OUTPUT}`
- Reference must remain byte-for-byte unchanged. No authoring until the visual gate is PASS.

## Page system

One portrait A4 section, w:pgSz 11906 by 16838 twips. Margins: top 1134,
bottom 1134, left 1701, right 850 twips, corresponding to 2 cm, 2 cm, 3 cm,
1.5 cm. Header and footer distances are 708 twips, approximately 1.25 cm.
Gutter 0. One column with space 708 twips; docGrid linePitch 360. No first-page
variant, header/footer parts, page-number fields, TOC, footnotes or content controls.
The final sectPr is cloned unchanged. The content width is 9355 twips or
5940425 EMU, approximately 16.5 cm. Do not introduce a standalone title page,
new section, generic page furniture, footer or page numbers.

## Source typography and paragraph components

Times New Roman is authoritative. Normal style id a has 12 pt text, color
1D2733, single line spacing, no space after. The template's real style ids,
theme and numbering are retained as bytes, including unused style definitions.
Direct formatting from the source patterns determines the applied appearance.

| Role | Stable source locator | Source formatting and use |
| --- | --- | --- |
| Institution | word/document.xml body paragraph 0 | Centered, 11 pt bold, after 2 pt |
| Department | body paragraph 1 | Centered, 11 pt, after 2 pt |
| Laboratory title | body paragraph 3 | Centered, 16 pt bold uppercase, after 6 pt |
| Topic | body paragraph 4 | Centered, 14 pt bold, after 10 pt |
| Team metadata | body paragraph 5 | Normal, 11 pt, four lines with source label alignment |
| Objective heading | body paragraph 7 | Normal 12 pt bold |
| Objective prose | body paragraph 9 | Normal 12 pt single line |
| Task heading | body paragraph 11 | Real style2, direct 12.5 pt bold, color1D2733, before12 pt after6 pt |
| Explanation | body paragraph24 | Normal12 pt, before5 pt after5 pt with source auto spacing |
| Subheading | body paragraph28 | Normal12 pt bold, keepNext, before5 pt after5 pt |
| Result | body paragraph35 | Normal10 pt, after6 pt |
| Figure | body paragraph40 | Centered drawing, before3 pt after5 pt |
| Caption | body paragraph69 | Normal10 pt, before3 pt after5 pt, keepLines |
| Closing heading | body paragraph77 | Real style2, keepNext, compact source spacing |
| Closing prose and oral answers | body paragraph78 | Normal10 pt, before3 pt after5 pt |

No default Word Title blue border or default blank-document styles are applied.
The title uses the actual source title component, whose style is Normal; this
is a deliberate template fidelity choice. Source task heading color is retained
and direct run properties; do not replace it with blank-document heading styles.
Keep heading with its following content, keep each PNG with its own caption,
and prevent a single text paragraph splitting unnecessarily. Cloning source
patterns may add keepNext/keepLines to support changed content without changing
font size or establishing a second visual system.

## Tables and recurring components

The source course table is body table0, three rows and two columns, 5000pct width,
source borders/margins and original cell paragraph/run properties. Rewrite course,
programme and lecturer text while retaining this component. The response table
is cloned from source body table1, Table Grid style a6, three columns and an
existing header and data row. Expand it to two data rows. Reuse real cell formats;
adjust its grid to 2500,2800,4055 twips for request, status and handler content.
Header repetition and non-splitting rows are permitted structural support for
longer report data. No fixed row heights or invented shaded callouts.

## Editable slot map and content flow

1. Clone institution/department, title/topic, course table and team metadata.
   Rewrite laboratory number6, precise topic, authors Aruzhan Zhaksybekova and
   Mukhamedjan Arman, group IT1-2404IS, date October12 2026 and variant11.
2. Clone objective components for aim, lab06 link and honest implementation/storage
   context. Introduce current27-test package verification separately from the
   October10 runtime evidence and its original18-test run.
3. Task1 form object: BookForm, Book record, Genre and GET-new source PNGs, then
   concise explanation of partial binding and the separate storage model.
4. Task2 template and POST source PNGs, then actual invalid-title/year/genre,
   simultaneous-error, success and refresh captures with concise observed results.
5. Task3 locale and both full resource bundle source PNGs, then actual Russian
   form/error capture. Preserve mandatory Russian evidence inside images.
6. Task4 custom annotation and validator source PNGs, then custom failure capture;
   API advice source PNGs and measured JSON404 capture.
7. Individual variant11 page advice source PNGs, HTML404 and common-model detail
   captures, and the measured two-response table. Each source/code capture appears
   exactly once; the common-model detail image remains Figure11.
8. Four-sentence conclusion comparing declarative/manual validation, then relevant
   concise oral answers cloned from the source compact prose component.

The old laboratory3 body is an editable replacement slot, including its old
paragraphs, tables, drawings and obsolete text. Its obsolete image parts and
image relationships are removed, as authorized by the root after the reference
review, while all non-content template parts remain byte-for-byte exact.
The new body references only33 requested PNGs:22 actual-source code captures and
11 actual runtime captures. No live editor was used; do not claim IDE screenshots.

## Code images and Word captions

The source PNGs are rendered from the actual project source bytes and have the
user-requested black/orange/cyan theme. File name, path and line-range headings
are absent from the PNGs. Write complete source path and correct original line
range only in each Word caption. Do not modify source or capture assets.
At original aspect ratio, code PNG width is the exact source content width,
approximately6.5 inches; runtime figures use the same width with at most4-inch
height to permit readable two-figure pages. Never stretch or crop evidence.

## Package preservation and fidelity gates

Only word/document.xml, word/_rels/document.xml.rels, docProps/core.xml and
[Content_Types].xml are editable package parts. Preserve every other reference
non-content part byte-for-byte, including styles, numbering, theme, fontTable,
settings, webSettings and app metadata. Add distinct Lab06 image relationships
and new media parts; preserve non-image relationship records unchanged. Keep
sectPr byte-equivalent. PNG bytes in the output must
match their supplied source bytes by SHA256. All22 capture raw/source hashes
must match the current files before authoring. No field refresh is required
because the reference has no fields.

Before delivery: inspect all reference and final pages at100percent original
resolution; compare title/task/body/table/figure patterns against the reference;
confirm exact section geometry; compare every preserve-only part and source
relationship; confirm all33 expected blips and images plus13 code groups; ensure
English narration and captions, required bilingual evidence, correct dates and
27-test current build proof. The final DOCX may be delivered only after the
latest render and structural review pass.

## Reproducible tracked support

The compact reports/lab06/report_template.docx contains the exact retained
non-content package parts and section geometry, with an empty body and no
obsolete Lab03 media. report_template_components.json contains the source-derived
paragraph and table patterns with old text cleared, retained reference hash,
template hash and nine-page PASS review metadata. The supported entry point is
reports/lab06/build_report.py --build. Default authoring uses those tracked files,
the image manifests and evidence/package_verification.json; ignored qa/reference
and target logs are not required after cloning. The small verification record is
derived from the actual October12 package log and includes all class counts,
source commit65ff8e8 and the raw-log SHA256. Its original log remains diagnostic
evidence, while this record supports stable reproduction.
"""
    if CONTRACT.exists():
        previous = CONTRACT.read_text(encoding="utf-8")
        for label in ["Reference visual gate", "Reference page count", "Reference render evidence"]:
            match = re.search(r"^" + re.escape(label) + r": .+$", previous, re.MULTILINE)
            if match:
                text = re.sub(r"^" + re.escape(label) + r": .+$", match.group(), text, flags=re.MULTILINE)
    CONTRACT.write_text(text, encoding="utf-8")
    print(json.dumps({"contract": str(CONTRACT), "baseline": str(BASELINE),
                      "reference_sha256": REFERENCE_HASH, "parts": len(inventory)}))


def first_run_properties(p):
    """Use actual text formatting, not the paragraph-mark formatting."""
    prop = p.find(tag("w", "r") + "/" + tag("w", "rPr"))
    if prop is None:
        prop = p.find(tag("w", "pPr") + "/" + tag("w", "rPr"))
    return deepcopy(prop) if prop is not None else element("w", "rPr")


def prepare_template():
    """Derive a compact tracked base after the retained reference visual gate."""
    verified_reference()
    if "Reference visual gate: PASS" not in CONTRACT.read_text(encoding="utf-8"):
        raise ValueError("Reference review is required before deriving the tracked template.")
    with ZipFile(REFERENCE) as src:
        items = [(deepcopy(i), src.read(i.filename)) for i in src.infolist()]
    parts = {i.filename: data for i, data in items}
    doc = ET.fromstring(parts["word/document.xml"])
    body = doc.find(tag("w", "body"))
    ps = [p for p in body if p.tag == tag("w", "p")]
    tables = [t for t in body if t.tag == tag("w", "tbl")]
    components = {"reference_sha256": REFERENCE_HASH,
                  "review": {"status": "PASS", "pages": 9, "date": "2026-10-12",
                             "render_evidence": "reports/lab06/qa/reference-render/page-1.png through page-9.png"},
                  "paragraphs": {}, "tables": []}
    for role, index in PROTOTYPES.items():
        clone = deepcopy(ps[index])
        for text in clone.findall(".//" + tag("w", "t")):
            text.text = ""
        for prop in clone.findall(".//" + tag("wp", "docPr")) + clone.findall(".//" + tag("pic", "cNvPr")):
            prop.set("name", "Reusable source image component")
            prop.attrib.pop("descr", None)
        components["paragraphs"][role] = ET.tostring(clone).decode()
    for table in tables[:2]:
        clone = deepcopy(table)
        for text in clone.findall(".//" + tag("w", "t")):
            text.text = ""
        components["tables"].append(ET.tostring(clone).decode())
    section = deepcopy(body.find(tag("w", "sectPr")))
    for child in list(body):
        body.remove(child)
    body.append(section)
    relationships = ET.fromstring(parts["word/_rels/document.xml.rels"])
    removed = []
    for relation in list(relationships):
        if relation.get("Type", "").endswith("/image"):
            removed.append(relation.get("Id"))
            relationships.remove(relation)
    types = ET.fromstring(parts["[Content_Types].xml"])
    for item in list(types):
        if item.get("PartName", "").startswith("/word/media/"):
            types.remove(item)
    changes = {"word/document.xml": ET.tostring(doc, xml_declaration=True, encoding="UTF-8", standalone=True),
               "word/_rels/document.xml.rels": ET.tostring(relationships, xml_declaration=True, encoding="UTF-8", standalone=True),
               "[Content_Types].xml": ET.tostring(types, xml_declaration=True, encoding="UTF-8", standalone=True)}
    core = ET.fromstring(parts["docProps/core.xml"])
    for name, text in [("title", "Lab 6 report template"), ("subject", "Retained Laboratory Work 3 document style")]:
        node = core.find("{http://purl.org/dc/elements/1.1/}" + name)
        if node is not None:
            node.text = text
    changes["docProps/core.xml"] = ET.tostring(core, xml_declaration=True, encoding="UTF-8", standalone=True)
    with ZipFile(TEMPLATE, "w") as out:
        for info, data in items:
            if not info.filename.startswith("word/media/"):
                out.writestr(info, changes.get(info.filename, data))
    components["template_sha256"] = sha256(TEMPLATE.read_bytes()).hexdigest()
    components["removed_source_image_relationships"] = removed
    components["preserved_parts"] = {name: sha256(data).hexdigest() for name, data in parts.items()
                                     if name not in EDITABLE_PARTS and not name.startswith("word/media/")}
    COMPONENTS.write_text(json.dumps(components, indent=2), encoding="utf-8")
    with ZipFile(TEMPLATE) as check:
        for name, digest in components["preserved_parts"].items():
            if sha256(check.read(name)).hexdigest() != digest:
                raise ValueError(f"Template part changed: {name}")
    print(json.dumps({"template": str(TEMPLATE), "bytes": TEMPLATE.stat().st_size,
                      "components": str(COMPONENTS), "removed_old_images": len(removed)}))


def record_package_verification():
    raw = BUILD_LOG.read_bytes()
    log = raw.decode("utf-8", errors="replace")
    pattern = r"Tests run: (\d+), Failures: (\d+), Errors: (\d+), Skipped: (\d+)"
    totals = re.findall(pattern, log)
    if not totals or "BUILD SUCCESS" not in log:
        raise ValueError("A successful package log is required to record verification.")
    current = dict(zip(["tests", "failures", "errors", "skipped"], map(int, totals[-1])))
    if current != {"tests": 27, "failures": 0, "errors": 0, "skipped": 0}:
        raise ValueError("The supplied package result is not the reviewed 27-test run.")
    classes = []
    for match in re.finditer(pattern + r", Time elapsed: ([\d.]+) s -- in ([\w.]+)", log):
        classes.append(dict(zip(["tests", "failures", "errors", "skipped"], map(int, match.groups()[:4])))
                       | {"class": match.group(6), "elapsed_seconds": float(match.group(5))})
    if sum(c["tests"] for c in classes) != current["tests"]:
        raise ValueError("Class counts differ from the package total.")
    finished = re.search(r"Finished at: (\S+)", log)
    data = {"date": "2026-10-12", "status": "BUILD SUCCESS", **current,
            "finished_at": finished.group(1) if finished else None,
            "source_commit": "65ff8e8", "raw_log": "target/memory-check/maven-package.log",
            "raw_log_sha256": sha256(raw).hexdigest(), "classes": classes}
    PACKAGE_VERIFICATION.write_text(json.dumps(data, indent=2), encoding="utf-8")
    print(json.dumps({"verification": str(PACKAGE_VERIFICATION), "tests": current["tests"]}))


class TemplateReport:
    def __init__(self):
        if REFERENCE.exists():
            verified_reference()
        components = json.loads(COMPONENTS.read_text(encoding="utf-8"))
        if components.get("review", {}).get("status") != "PASS":
            raise ValueError("Reference visual gate is pending; no DOCX authoring is allowed.")
        if sha256(TEMPLATE.read_bytes()).hexdigest() != components["template_sha256"]:
            raise ValueError("Tracked template has changed; repeat template distillation.")
        with ZipFile(TEMPLATE) as z:
            self.items = [(deepcopy(info), z.read(info.filename)) for info in z.infolist()]
            self.original_parts = {info.filename: data for info, data in self.items}
        self.document = ET.fromstring(self.original_parts["word/document.xml"])
        self.body = self.document.find(tag("w", "body"))
        self.source_paragraphs = {role: ET.fromstring(xml.encode()) for role, xml in components["paragraphs"].items()}
        self.source_tables = [ET.fromstring(xml.encode()) for xml in components["tables"]]
        self.section = deepcopy(self.body.find(tag("w", "sectPr")))
        self.relations = ET.fromstring(self.original_parts["word/_rels/document.xml.rels"])
        self.new_parts = {}
        self.image_count = 0
        self.code_count = 0
        self.figure_count = 0
        self.listing_count = 0
        self.content_width = 5940425
        for child in list(self.body):
            self.body.remove(child)
        self.code_manifest = json.loads((EVIDENCE / "code_manifest.json").read_text(encoding="utf-8-sig"))
        evidence_manifest = json.loads((EVIDENCE / "manifest.json").read_text(encoding="utf-8-sig"))
        self.evidence_manifest = evidence_manifest
        self.runtime = {Path(r["path"]).name: english_screenshot(r) for r in evidence_manifest["screenshots"]}

    def p(self, role: str, text: str, keep_next=False):
        p = deepcopy(self.source_paragraphs[role])
        ppr = p.find(tag("w", "pPr"))
        rpr = first_run_properties(p)
        for child in list(p):
            if child is not ppr:
                p.remove(child)
        if ppr is None:
            ppr = element("w", "pPr")
            p.insert(0, ppr)
        for old in list(ppr):
            if old.tag == tag("w", "sectPr"):
                ppr.remove(old)
        if keep_next and ppr.find(tag("w", "keepNext")) is None:
            ppr.append(element("w", "keepNext"))
        run = element("w", "r")
        run.append(rpr)
        for i, line in enumerate(text.split("\n")):
            if i:
                run.append(element("w", "br"))
            t = element("w", "t")
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            t.text = line
            run.append(t)
        p.append(run)
        self.body.append(p)
        return p

    @staticmethod
    def rewrite_cell(cell, text: str):
        ps = cell.findall(tag("w", "p"))
        p = ps[0] if ps else element("w", "p")
        if not ps:
            cell.append(p)
        ppr = p.find(tag("w", "pPr"))
        rpr = first_run_properties(p)
        for child in list(p):
            if child is not ppr:
                p.remove(child)
        for extra in ps[1:]:
            cell.remove(extra)
        run = element("w", "r")
        run.append(rpr)
        for i, line in enumerate(text.split("\n")):
            if i:
                run.append(element("w", "br"))
            t = element("w", "t")
            t.text = line
            run.append(t)
        p.append(run)

    def course_table(self):
        table = deepcopy(self.source_tables[0])
        rows = table.findall(tag("w", "tr"))
        values = [
            ("Discipline:", "RWPSF 3305 - Web Application Development with Spring Framework"),
            ("Educational programme:", "6B06105 - Information Systems, year 3, semester 6"),
            ("Lecturer:", "Yemberdiyeva Aknur, Senior Lector"),
        ]
        for row, pair in zip(rows, values):
            for cell, text in zip(row.findall(tag("w", "tc")), pair):
                self.rewrite_cell(cell, text)
        self.body.append(table)

    def image(self, file: Path, caption: str, runtime=False):
        data = file.read_bytes()
        with Image.open(file) as im:
            w, h = im.size
        cx = self.content_width
        if runtime:
            cx = min(cx, round(4 * 914400 * w / h))
        cy = round(cx * h / w)
        if cy > 8.7 * 914400:
            raise ValueError(f"Image is too tall for the source page system: {file}")
        self.image_count += 1
        rid = f"rIdLab06Image{self.image_count:03d}"
        target = f"media/lab06_image_{self.image_count:03d}.png"
        rel = ET.SubElement(self.relations, tag("rel", "Relationship"))
        rel.set("Id", rid)
        rel.set("Type", NS["r"] + "/image")
        rel.set("Target", target)
        self.new_parts["word/" + target] = data
        p = deepcopy(self.source_paragraphs["image"])
        ppr = p.find(tag("w", "pPr"))
        ppr.append(element("w", "keepNext"))
        for t in p.findall(".//" + tag("w", "t")):
            t.text = ""
        for extent in p.findall(".//" + tag("wp", "extent")) + p.findall(".//" + tag("a", "ext")):
            extent.set("cx", str(cx))
            extent.set("cy", str(cy))
        for blip in p.findall(".//" + tag("a", "blip")):
            blip.set(tag("r", "embed"), rid)
        for prop in p.findall(".//" + tag("wp", "docPr")) + p.findall(".//" + tag("pic", "cNvPr")):
            prop.set("id", str(1000 + self.image_count))
            prop.set("name", f"Lab06 evidence {self.image_count}")
            prop.set("descr", caption)
        self.body.append(p)
        self.p("caption", caption)

    def code(self, key: str, title: str):
        records = self.code_manifest.get(key, [])
        if not records:
            raise ValueError(f"Missing actual source captures for {key}")
        self.listing_count += 1
        self.p("subheading", title, keep_next=True)
        for i, record in enumerate(records, 1):
            path = ROOT / record["source"]
            if sha256(path.read_bytes()).hexdigest() != record["source_sha256"]:
                raise ValueError(f"Capture no longer matches actual source bytes: {path}")
            desc = re.sub(r" part \d+ of \d+$", "", record.get("caption", title))
            part = f" part {i} of {len(records)}" if len(records) > 1 else ""
            caption = (f"Listing {self.listing_count}{part}. {desc}. "
                       f"Source: {record['source']}, lines {record['start_line']} to {record['end_line']}.")
            self.image(EVIDENCE / record["file"], caption)
            self.code_count += 1

    def result(self, filename: str):
        record = self.runtime[filename]
        self.figure_count += 1
        self.image(EVIDENCE / filename,
                   f"Figure {self.figure_count}. {record['caption']}.", runtime=True)
        self.p("result", record.get("details", ""))

    def response_table(self):
        table = deepcopy(self.source_tables[1])
        rows = table.findall(tag("w", "tr"))
        while len(rows) < 3:
            row = deepcopy(rows[-1])
            table.append(row)
            rows.append(row)
        grid = [2500, 2800, 4055]
        for col, width in zip(table.find(tag("w", "tblGrid")), grid):
            col.set(tag("w", "w"), str(width))
        values = [
            ["Request", "Status and Content Type", "Handler and result"],
            ["GET /books/999", "404\ntext/html;charset=UTF-8",
             "PageExceptionHandler.bookNotFound\nLocalized error/not-found HTML page for identifier 999."],
            ["GET /api/books/999", "404\napplication/problem+json",
             "ApiExceptionHandler.bookNotFound\nProblemDetail with status=404 and instance=/api/books/999."],
        ]
        for i, (row, texts) in enumerate(zip(rows, values)):
            trpr = row.find(tag("w", "trPr"))
            if trpr is None:
                trpr = element("w", "trPr")
                row.insert(0, trpr)
            trpr.append(element("w", "cantSplit"))
            if i == 0:
                trpr.append(element("w", "tblHeader"))
            for cell, text, width in zip(row.findall(tag("w", "tc")), texts, grid):
                tcw = cell.find(tag("w", "tcPr") + "/" + tag("w", "tcW"))
                if tcw is not None:
                    tcw.set(tag("w", "w"), str(width))
                self.rewrite_cell(cell, text)
        self.body.append(table)

    def save(self):
        self.body.append(self.section)
        paragraphs = self.document.findall(".//" + tag("w", "p"))
        all_text = "\n".join("".join(p.itertext()) for p in paragraphs)
        if re.search(r"[\u0400-\u04ff]", all_text):
            raise ValueError("Report narration must be English; Russian remains only in required evidence PNGs.")
        if (self.image_count, self.code_count, self.figure_count, self.listing_count) != (33, 22, 11, 13):
            raise ValueError("Expected 33 images, 22 source captures, 11 results and 13 code groups.")
        changes = {
            "word/document.xml": ET.tostring(self.document, xml_declaration=True, encoding="UTF-8", standalone=True),
            "word/_rels/document.xml.rels": ET.tostring(self.relations, xml_declaration=True, encoding="UTF-8", standalone=True),
        }
        core = ET.fromstring(self.original_parts["docProps/core.xml"])
        core_ns = {"dc": "http://purl.org/dc/elements/1.1/",
                   "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
                   "dcterms": "http://purl.org/dc/terms/"}
        for prefix, local, value in [
            ("dc", "title", "Laboratory Work 6 Forms and Validation"),
            ("dc", "subject", "RWPSF 3305 Web Application Development with Spring Framework"),
            ("dc", "creator", "Aruzhan Zhaksybekova; Mukhamedjan Arman"),
            ("cp", "lastModifiedBy", "Aruzhan Zhaksybekova; Mukhamedjan Arman"),
            ("dcterms", "modified", "2026-10-12T00:00:00Z"),
        ]:
            node = core.find("{" + core_ns[prefix] + "}" + local)
            if node is None:
                node = ET.SubElement(core, "{" + core_ns[prefix] + "}" + local)
            node.text = value
        changes["docProps/core.xml"] = ET.tostring(core, xml_declaration=True, encoding="UTF-8", standalone=True)
        content_types = ET.fromstring(self.original_parts["[Content_Types].xml"])
        if not any(x.get("Extension") == "png" for x in content_types):
            png = ET.SubElement(content_types, "{http://schemas.openxmlformats.org/package/2006/content-types}Default")
            png.set("Extension", "png")
            png.set("ContentType", "image/png")
            changes["[Content_Types].xml"] = ET.tostring(content_types, xml_declaration=True, encoding="UTF-8", standalone=True)
        with ZipFile(OUTPUT, "w") as out:
            for info, data in self.items:
                out.writestr(info, changes.get(info.filename, data))
            for name, data in self.new_parts.items():
                out.writestr(name, data)
        with ZipFile(OUTPUT) as check:
            for name, data in self.original_parts.items():
                if name not in EDITABLE_PARTS and check.read(name) != data:
                    raise ValueError(f"Preserve-only package part changed: {name}")
            for name, data in self.new_parts.items():
                if check.read(name) != data:
                    raise ValueError(f"Evidence bytes changed: {name}")
        if REFERENCE.exists():
            verified_reference()
        print(json.dumps({"output": str(OUTPUT), "images": self.image_count,
                          "code_images": self.code_count, "runtime_images": self.figure_count,
                          "reference_preserved": True, "opaque_parts_preserved": True}))


def latest_build_result() -> str:
    verification = json.loads(PACKAGE_VERIFICATION.read_text(encoding="utf-8"))
    tests, failures, errors, skipped = [verification[k] for k in ["tests", "failures", "errors", "skipped"]]
    if verification["status"] != "BUILD SUCCESS" or failures or errors:
        raise ValueError("Package verification record contains failures.")
    return (f"Current package verification on October 12, 2026: BUILD SUCCESS. "
            f"Tests run: {tests}, failures: {failures}, errors: {errors}, skipped: {skipped}. "
            "Application evidence below was captured "
            "on October 10, 2026; that original verification run had 18 passing tests.")


def build():
    r = TemplateReport()
    r.p("institution", "JSC International Information Technology University")
    r.p("department", "Faculty of Business, Media and Management - Department of Information Systems")
    r.p("title", "LABORATORY WORK No. 6")
    r.p("topic", "Forms and validation\nThymeleaf binding, Bean Validation, localization and centralized exception handling")
    r.course_table()
    r.p("metadata", "Performed by:                      Aruzhan Zhaksybekova & Mukhamedjan Arman\n"
                      "Group:                                  IT1-2404IS\n"
                      "Date:                                     October 12, 2026\n"
                      "Variant:                                 11")
    r.p("objective_heading", "Objective of the work", keep_next=True)
    r.p("body", "The objective is to implement a complete book form with Thymeleaf binding, Jakarta Bean "
                "Validation, localized messages, a custom constraint and separate HTML/API exception responses. "
                "Invalid input must retain the entered values; valid input must finish with Post/Redirect/Get.")
    r.p("compact_body", "Repository branch: " + REPO_LINK)
    r.p("compact_body", latest_build_result())

    r.p("task_heading", "Task 1 The form object and the GET handler", keep_next=True)
    r.p("explanation", "We added the Thymeleaf and validation starters, restored the missing book catalogue "
                      "in the supplied Laboratory Work 4 project, and placed the module in kz.iitu.springlab.lab6.")
    r.code("BookForm.java", "BookForm and its validation rules")
    r.code("Book.java", "The separate stored Book record")
    r.code("Genre.java", "The available genres")
    r.code("get_new", "GET books new and the initial model")
    r.p("result", "BookForm has setters so binding can populate it partially and retain rejected values. "
                  "Integer distinguishes an empty year from a numeric year; @NotNull checks presence and "
                  "@Min/@Max check the 1450-2100 range. Book contains the service-assigned identifier and "
                  "stored values. GET /books/new supplies both form and genres. BookService stores records "
                  "in process memory, so restarting the application restores the three initial books.")

    r.p("task_heading", "Task 2 The template and the POST handler", keep_next=True)
    r.code("form.html", "The complete Thymeleaf form")
    r.code("post", "Validation and Post Redirect Get in the POST handler")
    r.p("result", "th:object selects form; th:field generates the input name, id and current value. "
                  "th:errors and #fields display validation messages, while th:errorclass marks invalid fields. "
                  "BindingResult follows @Valid BookForm immediately. The error branch restores genres "
                  "and returns books/form; a valid submission saves the book and redirects with a flash message.")
    for file in ["01-empty-title-en.png", "02-year-out-of-range-en.png", "03-no-genre-en.png",
                 "06-two-errors-ru.png", "07-success-ru.png", "08-refresh-no-duplicate.png"]:
        r.result(file)

    r.p("task_heading", "Task 3 Two languages and localized messages", keep_next=True)
    r.code("LocaleConfiguration.java", "Locale selection through the lang parameter")
    r.code("messages.properties", "The complete default English resource bundle")
    r.code("messages_ru.properties", "The complete Russian resource bundle")
    r.p("result", "SessionLocaleResolver keeps the selected language in the HTTP session. "
                  "LocaleChangeInterceptor handles ?lang=en and ?lang=ru. Both captions and constraints "
                  "resolve resource keys through MessageSource; th:errors then displays the localized errors. "
                  "The default messages.properties file enables message-source configuration and supplies "
                  "fallback values. Figure 1 and the following capture show the same title violation in both languages.")
    r.result("04-empty-title-ru.png")

    r.p("task_heading", "Task 4 A custom constraint and exception advice", keep_next=True)
    r.code("NoDigits.java", "The custom NoDigits annotation")
    r.code("NoDigitsValidator.java", "The NoDigits validator implementation")
    r.p("result", "The annotation declares message, groups and payload and selects NoDigitsValidator. "
                  "The validator accepts null because @NotBlank separately checks required input. "
                  "Entering Bloch 2018 fails the digit check and resolves book.author.nodigits from the selected bundle.")
    r.result("05-custom-constraint-ru.png")
    r.code("ApiExceptionHandler.java", "ProblemDetail handling for REST controllers")
    r.p("result", "BookService throws BookNotFoundException for an unknown identifier. "
                  "ApiExceptionHandler is scoped to @RestController classes and returns status 404 "
                  "with an application/problem+json body. The HTML handler is also shown in the individual assignment below.")
    r.result("10-api-problem-detail.png")

    r.p("task_heading", "Individual assignment Variant 11", keep_next=True)
    r.code("PageExceptionHandler.java", "Shared model attribute and HTML exception handling")
    r.p("result", "Variant 11 adds @ControllerAdvice with @ModelAttribute and an exception handler. "
                  "PageExceptionHandler targets BookPageController. Its libraryName method supplies the "
                  "localized shared heading to catalogue pages, and bookNotFound returns error/not-found "
                  "with HTTP 404. The shared value appears on the form, catalogue, error and detail pages.")
    r.result("09-page-not-found.png")
    r.result("11-variant11-common-model.png")
    r.p("subheading", "Two responses for the same missing identifier", keep_next=True)
    r.response_table()

    r.p("compact_heading", "Conclusions", keep_next=True)
    r.p("compact_body", "Declarative validation places input rules in BookForm instead of repeating manual "
                        "checks in the controller. BindingResult lets the form display all violations while "
                        "retaining the submitted values. Resource bundles separate text from code and keep "
                        "labels and validation messages in the selected language. Centralized advice gives "
                        "HTML and API clients consistent error responses, while Post/Redirect/Get prevents "
                        "a result-page refresh from repeating the POST.")
    r.p("compact_heading", "Answers for the demonstration", keep_next=True)
    answers = [
        "A separate form object limits the accepted fields and holds input rules, while the stored record contains its identifier and storage values.",
        "th:object selects the form object. After it, *{title} refers to that object's title; th:field generates name, id and a retained value.",
        "GET must supply an empty form object so Thymeleaf can resolve its binding expressions, together with the available genres.",
        "Without @Valid, Bean Validation constraints are not invoked by this handler, although binding can still report conversion errors.",
        "BindingResult must immediately follow the bound object. It collects field errors, allowing an ordinary error branch without throwing a validation exception.",
        "The error branch restores genres because the new request model must again contain the options required by the template.",
        "@NotNull rejects null, @NotEmpty also rejects an empty value, and @NotBlank rejects empty or whitespace-only strings.",
        "The annotation's message key is resolved from the locale's bundle and displayed by th:errors. A field constraint produces a field error; a class-level constraint can produce a global error.",
        "A constraint declares message for the error text, groups for selecting validation contexts and payload for extra metadata. NoDigits accepts null because required input is checked separately.",
        "Advice is scoped by controller type. Unscoped advice could compete to handle the same exception with the wrong response format; local handlers, advice order and match specificity affect selection.",
        "th:text escapes HTML. th:utext is unsuitable for untrusted user input because it can expose an XSS vulnerability.",
    ]
    for i, answer in enumerate(answers, 1):
        r.p("compact_body", f"{i}. {answer}")
    r.save()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--prepare", action="store_true")
    parser.add_argument("--build", action="store_true")
    parser.add_argument("--prepare-template", action="store_true")
    parser.add_argument("--record-verification", action="store_true")
    args = parser.parse_args()
    if args.prepare:
        prepare_contract()
    if args.prepare_template:
        prepare_template()
    if args.record_verification:
        record_package_verification()
    if args.build:
        build()
    if not any([args.prepare, args.build, args.prepare_template, args.record_verification]):
        print("Use --prepare for the contract; --build only after the reference visual gate passes.")


if __name__ == "__main__":
    main()
