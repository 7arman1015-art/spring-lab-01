"""Build the Lab 6 report from the current source files and genuine evidence.

Run with the bundled Codex Python after the document-operation marker succeeds.
This script deliberately does not run Maven, change Git state, fabricate pictures,
or render the resulting DOCX. Its output must pass render and visual QA separately.
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[2]
REPORT_DIR = ROOT / "reports" / "lab06"
DEFAULT_MANIFEST = REPORT_DIR / "evidence" / "manifest.json"
REPO_LINK = "https://github.com/7arman1015-art/spring-lab-01/tree/lab06"


def source(name: str) -> tuple[str, str]:
    matches = list((ROOT / "src" / "main").rglob(name))
    if len(matches) != 1:
        raise RuntimeError(f"Expected one source file {name}; found {matches}")
    path = matches[0]
    return path.relative_to(ROOT).as_posix(), path.read_text(encoding="utf-8-sig")


def field(paragraph, instruction: str, display: str = "1") -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run = paragraph.add_run()
    run._r.append(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    paragraph.add_run()._r.append(instr)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    paragraph.add_run()._r.append(separate)
    paragraph.add_run(display)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(end)


def hyperlink(paragraph, label: str, target: str) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    prop = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    prop.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    prop.append(underline)
    run.append(prop)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def configure(doc) -> None:
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(2.0), Cm(2.0)
    sec.footer_distance = Cm(0.8)
    sec.different_first_page_header_footer = True
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        if name.startswith("Heading"):
            style.font.size = Pt(14 if name == "Heading 1" else 12)
            style.font.bold = True
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.keep_with_next = True
    doc.styles["Title"].font.size = Pt(21)
    doc.styles["Subtitle"].font.size = Pt(13)
    caption = doc.styles["Caption"]
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.paragraph_format.space_after = Pt(7)
    code = doc.styles.add_style("Code Listing", 1)
    code.font.name = "Consolas"
    code.font.size = Pt(8)
    code.font.color.rgb = RGBColor(0, 0, 0)
    code.paragraph_format.line_spacing = Pt(9.3)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.keep_with_next = False
    code.paragraph_format.widow_control = False
    # Word's packaged template includes a blue Title border and theme font slots.
    # Remove them rather than letting LibreOffice/Word override the chosen fonts.
    chosen_styles = [normal, code, caption] + [
        doc.styles[name] for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3")
    ]
    for st in chosen_styles:
        for border in st.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
        fonts = st.element.get_or_add_rPr().get_or_add_rFonts()
        for theme in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme", "csTheme"):
            fonts.attrib.pop(qn("w:" + theme), None)
        for script in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + script), st.font.name)
        for color in st.element.xpath("./w:rPr/w:color"):
            for attr in ("themeColor", "themeTint", "themeShade"):
                color.attrib.pop(qn("w:" + attr), None)
            color.set(qn("w:val"), "000000")
        if st.font.size is not None:
            props = st.element.get_or_add_rPr()
            size_cs = props.find(qn("w:szCs"))
            if size_cs is None:
                size_cs = OxmlElement("w:szCs")
                props.append(size_cs)
            size_cs.set(qn("w:val"), str(round(st.font.size.pt * 2)))
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(footer, "PAGE")
    for run in footer.runs:
        run.font.size = Pt(9)
        run.font.name = "Times New Roman"
    doc.core_properties.title = "Laboratory Work 6 Forms and Validation"
    doc.core_properties.subject = "RWPSF 3305 Web Application Development with Spring Framework"
    doc.core_properties.author = "Aruzhan Zhaksybekova; Mukhamedjan Arman"
    doc.core_properties.keywords = "Spring MVC, Thymeleaf, Bean Validation, i18n, ControllerAdvice, variant 11"
    doc.core_properties.language = "en-US"


def paragraph(doc, text: str, **kwargs):
    p = doc.add_paragraph(text, **kwargs)
    if p.style.name in ("Title", "Subtitle"):
        for border in p._p.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    return p


def heading(doc, text: str, level: int = 1):
    return doc.add_heading(text, level=level)


def table(doc, headers: list[str], rows: list[list[str]], widths: list[float]):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    tblpr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement("w:" + edge)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")
        borders.append(element)
    tblpr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge, value in (("top", "70"), ("bottom", "70"), ("left", "90"), ("right", "90")):
        item = OxmlElement("w:" + edge)
        item.set(qn("w:w"), value)
        item.set(qn("w:type"), "dxa")
        margins.append(item)
    tblpr.append(margins)
    header_repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(header_repeat)
    for i, title in enumerate(headers):
        t.rows[0].cells[i].text = title
    for row in rows:
        cells = t.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)
    for row_index, row in enumerate(t.rows):
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for col_index, cell in enumerate(row.cells):
            cell.width = Cm(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E7EBF0" if row_index == 0 else "FFFFFF")
            cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.bold = row_index == 0
                    r.font.color.rgb = RGBColor(0, 0, 0)
    paragraph(doc, "").paragraph_format.space_after = Pt(0)
    return t


def listing(doc, title: str, file_name: str, number: int,
            code_manifest: dict, evidence_dir: Path, key: str | None = None):
    """Insert real captures of source code, never a plaintext substitute."""
    rel, _ = source(file_name)
    records = code_manifest.get(key or file_name, [])
    if not isinstance(records, list) or not records:
        raise ValueError(f"Missing code screenshots for {key or file_name}")
    heading(doc, title, level=2)
    cap = paragraph(doc, f"Listing {number}. {rel}", style="Caption")
    cap.paragraph_format.keep_with_next = True
    for i, record in enumerate(records, 1):
        image_path = Path(record["file"])
        if not image_path.is_absolute():
            image_path = evidence_dir / image_path
        if not image_path.is_file():
            raise FileNotFoundError(image_path)
        start, end = int(record["start_line"]), int(record["end_line"])
        if start < 1 or end < start:
            raise ValueError(f"Invalid original source line range for {image_path}")
        caption = english_text(record.get("caption", title))
        caption = re.sub(r" part \d+ of \d+$", "", caption)
        with Image.open(image_path) as im:
            width_px, height_px = im.size
        if 6.5 * height_px / width_px > 8.6:
            raise ValueError(f"Split this code screenshot into shorter sections: {image_path}")
        p = paragraph(doc, "")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_with_next = True
        pic = p.add_run().add_picture(str(image_path), width=Inches(6.5))
        pic._inline.docPr.set("descr", f"{rel}, source lines {start} to {end}: {caption}")
        part = f" part {i} of {len(records)}" if len(records) > 1 else ""
        cap = paragraph(doc, f"Listing {number}{part}. {caption}. Source lines {start} to {end}.", style="Caption")
        cap.paragraph_format.keep_with_next = False


def evidence_image(doc, record: dict, evidence_dir: Path, number: int):
    image_path = Path(record["path"])
    if not image_path.is_absolute():
        image_path = evidence_dir / image_path
    if not image_path.is_file():
        raise FileNotFoundError(image_path)
    with Image.open(image_path) as im:
        width_px, height_px = im.size
    width = min(6.50, 4.00 * width_px / height_px)
    p = paragraph(doc, "")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_with_next = True
    pic = p.add_run().add_picture(str(image_path), width=Inches(width))
    description = record.get("details", record.get("caption", image_path.stem))
    pic._inline.docPr.set("descr", str(description))
    cap = paragraph(doc, f"Figure {number}. {record.get('caption', image_path.stem)}", style="Caption")
    cap.paragraph_format.keep_with_next = bool(record.get("details"))
    if record.get("details"):
        paragraph(doc, record["details"])


def validate_manifest(manifest: dict):
    screenshots = manifest.get("screenshots", [])
    responses = manifest.get("responses", [])
    if len(screenshots) < 4:
        raise ValueError("The report needs at least four genuine screenshots.")
    if len(responses) < 2:
        raise ValueError("The report needs the measured HTML and API error responses.")
    if not isinstance(manifest.get("build"), dict):
        raise ValueError("Include build/test metadata; do not invent a passed result.")
    return screenshots, responses


LEGACY_ENGLISH = {
    "Ветка lab06 опубликована в GitHub. Код, доказательства и отчёт сохранены в репозитории.":
        "The lab06 branch is published on GitHub. The code, evidence and report are stored in the repository.",
    "Локальная ветка lab06 подготовлена. Отправка в GitHub ожидает подтверждения.":
        "The local lab06 branch is ready. Publication on GitHub is awaiting confirmation.",
    "Maven 3.9.16: mvn -B package (прямой запуск Maven из кэша Wrapper)":
        "Maven 3.9.16: mvn -B package, run directly from the Maven Wrapper cache.",
    "Отклонённые формы": "Rejected submissions",
    "Шесть некорректных отправок не добавили ни одной книги.":
        "Six invalid submissions did not create any books.",
    "Вариант 11": "Variant 11",
    "libraryName отображается в общей шапке формы, каталога, деталей и страницы ошибки.":
        "libraryName appears in the shared header of the form, catalogue, detail page and error page.",
    "Ответы 404": "404 responses",
    "HTML: 404 text/html; API: 404 application/problem+json. Содержимое и заголовки сохранены отдельно.":
        "HTML: 404 text/html. API: 404 application/problem+json. The response bodies and headers were also saved.",
    "Браузер": "Browser",
    "Реальные формы отправлены через headless Google Chrome. Ошибок pageerror не обнаружено.":
        "The forms were submitted in a real headless Google Chrome session. No pageerror events were detected.",
    "HTML error/not-found; «Книга не найдена», идентификатор 999.":
        "HTML view error/not-found with a localized Book not found message for identifier 999.",
    "JSON ProblemDetail с status=404 и instance=/api/books/999.":
        "JSON ProblemDetail with status=404 and instance=/api/books/999.",
    "Проверено в реальном Chrome: POST /books → HTTP 302 (Location: http://localhost:8080/books) → GET /books. Создана одна книга; после reload flash-сообщение отсутствует и число книг остаётся 4.":
        "Verified in real Chrome: POST /books returned HTTP 302 with Location: http://localhost:8080/books, "
        "followed by GET /books. One book was created. After reloading the page, the flash message disappeared "
        "and the total remained four books.",
}

# The captured Russian UI and messages_ru source remain bilingual evidence.
# All report narration, including legacy evidence metadata, is English.
EVIDENCE_ENGLISH = {
    "01-empty-title-en.png": (
        "An empty title produces an English validation message",
        "POST /books returned HTTP 200 with the same form. The title error appears next to the field "
        "and in the summary. The author, year, genre and availability checkbox retain their values."),
    "02-year-out-of-range-en.png": (
        "The year 1200 is rejected",
        "@Min rejects the year 1200. The submitted values remain in the form, and the year input has a red border."),
    "03-no-genre-en.png": (
        "The genre is not selected",
        "@NotNull requires a genre. The title, author, year and availability flag are retained; "
        "the controller restores the genre options in the model."),
    "04-empty-title-ru.png": (
        "The same form and validation message in Russian",
        "After ?lang=ru, the labels, genres, buttons, shared heading and validation message appear in Russian. "
        "The selected locale remains in the HTTP session for the POST request."),
    "05-custom-constraint-ru.png": (
        "The custom NoDigits constraint rejects an author containing digits",
        "NoDigitsValidator rejects Bloch 2018. The Russian error message is resolved from "
        "book.author.nodigits in the Russian resource bundle."),
    "06-two-errors-ru.png": (
        "Two field errors are shown together",
        "The year 1200 and the missing genre are checked in the same submission. "
        "#fields.allErrors() lists both violations; each message also appears beside its field."),
    "07-success-ru.png": (
        "A valid submission creates a book",
        "The valid POST returned HTTP 302 and redirected to /books. The following GET shows a one-time "
        "flash message and the new Effective Java book with its selected genre and availability."),
    "08-refresh-no-duplicate.png": (
        "Reloading the result does not create a duplicate",
        "After reloading GET /books, the flash message is gone. The book count remains unchanged, "
        "and Effective Java appears exactly once."),
    "09-page-not-found.png": (
        "An HTML error page for a missing book",
        "GET /books/999 returned HTTP 404 with the error/not-found HTML view. "
        "PageExceptionHandler also supplies the shared libraryName attribute."),
    "10-api-problem-detail.png": (
        "A ProblemDetail response for the API",
        "GET /api/books/999 returned HTTP 404 with Content-Type application/problem+json "
        "and the type, title, status, detail and instance fields. ApiExceptionHandler handled the request."),
    "11-variant11-common-model.png": (
        "Variant 11 supplies the shared attribute on the book detail page",
        "The detail page displays libraryName supplied by @ModelAttribute in PageExceptionHandler. "
        "The same attribute is used by the form, catalogue and HTML error page."),
}


def english_text(value) -> str:
    text = str(value)
    translated = LEGACY_ENGLISH.get(text, text)
    if re.search(r"[\u0400-\u04ff]", translated):
        raise ValueError(f"English report metadata is required: {text}")
    return translated


def english_screenshot(record: dict) -> dict:
    result = dict(record)
    caption = str(record.get("caption", ""))
    details = str(record.get("details", ""))
    if re.search(r"[\u0400-\u04ff]", caption + details):
        known = EVIDENCE_ENGLISH.get(Path(record["path"]).name)
        if known is None:
            raise ValueError(f"Translate this screenshot metadata into English: {record['path']}")
        result["caption"], result["details"] = known
        urls = re.findall(r"https?://[^\s]+", details)
        if urls:
            result["details"] += " URL: " + urls[-1]
    else:
        result["caption"] = english_text(caption)
        result["details"] = english_text(details)
    return result


def build(manifest_path: Path, output: Path, code_manifest_path: Path | None = None):
    manifest = json.loads(manifest_path.read_text(encoding="utf-8-sig"))
    screenshots, responses = validate_manifest(manifest)
    screenshots = [english_screenshot(record) for record in screenshots]
    code_manifest_path = code_manifest_path or manifest_path.parent / "code_manifest.json"
    code_manifest = json.loads(code_manifest_path.read_text(encoding="utf-8-sig"))
    if not isinstance(code_manifest, dict):
        raise ValueError("code_manifest.json must map source filename or method key to screenshot records.")
    doc = Document()
    configure(doc)

    def add_listing(title: str, name: str, number: int, key: str | None = None):
        listing(doc, title, name, number, code_manifest, manifest_path.parent, key)

    for text in ("International Information Technology University",
                 "Faculty of Business Media and Management",
                 "Department of Information Systems"):
        p = paragraph(doc, text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph(doc, "").paragraph_format.space_after = Pt(30)
    p = paragraph(doc, "Laboratory Work 6", style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = paragraph(doc, "Forms and Validation", style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = paragraph(doc, "Thymeleaf  Bean Validation  Localization  Exception Handling")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph(doc, "").paragraph_format.space_after = Pt(28)
    for text in ("Course RWPSF 3305 Web Application Development with Spring Framework",
                 "Educational Programme 6B06105 Information Systems",
                 "Team members Aruzhan Zhaksybekova and Mukhamedjan Arman",
                 "Group IT1-2404IS",
                 "Individual Variant 11",
                 "Lecturer Yemberdiyeva Aknur"):
        p = paragraph(doc, text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph(doc, "").paragraph_format.space_after = Pt(25)
    p = paragraph(doc, "Almaty 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    heading(doc, "Aim and implemented result")
    paragraph(doc, "We implemented a book creation form with Spring MVC and Thymeleaf. "
                   "The form binds request parameters to BookForm, validates the fields with Jakarta Bean Validation, "
                   "displays messages in English and Russian, and saves valid input through Post/Redirect/Get. "
                   "A missing book produces an HTML error page or a JSON ProblemDetail response, both with status 404.")
    paragraph(doc, "The aim is to combine form binding, declarative validation and centralized exception handling "
                   "while preserving the submitted values when an invalid form is displayed again.")
    paragraph(doc, "The supplied project contained Laboratory Work 4. We restored the missing book catalogue "
                   "needed for Laboratory Work 6 in the kz.iitu.springlab.lab6 namespace. "
                   "BookService stores books in process memory; restarting the application restores the three initial records.")
    p = paragraph(doc, "Repository branch lab06: ")
    hyperlink(p, REPO_LINK, REPO_LINK)
    publication = manifest.get("publication", {})
    if publication.get("state"):
        paragraph(doc, "Repository status: " + english_text(publication["state"]))
    if publication.get("commit"):
        paragraph(doc, "Initial implementation commit: " + english_text(publication["commit"]))

    heading(doc, "Implementation steps")
    steps = [
        "Added spring-boot-starter-thymeleaf and spring-boot-starter-validation. "
        "Created BookForm, the Genre enumeration, the separate Book record and the book storage service.",
        "Added GET /books/new with an empty form object and the genre list. "
        "Bound the template with th:object and th:field, and displayed errors using th:errors, th:errorclass and #fields.",
        "Implemented POST /books with BindingResult immediately after the @Valid BookForm parameter. "
        "Invalid input returns the same form with the genre options restored. "
        "Valid input creates a book, adds a flash message and redirects to /books.",
        "Moved labels, error messages, genre names and the success message into messages.properties "
        "and messages_ru.properties. Configured SessionLocaleResolver and LocaleChangeInterceptor "
        "with the lang parameter and English as the default locale.",
        "Created @NoDigits and NoDigitsValidator for the author field. The validator accepts null; "
        "@NotBlank separately enforces required input. Added separate advice for HTML pages and the API.",
        "Implemented Variant 11 with @ModelAttribute libraryName in PageExceptionHandler. "
        "The localized library name is supplied automatically to the catalogue page models.",
        "Verified invalid and valid submissions in the browser, both locales, both 404 responses "
        "and a reload after a successful save. The following sections contain the measured results and screenshots.",
    ]
    for i, text in enumerate(steps, 1):
        paragraph(doc, f"{i}. {text}")

    heading(doc, "Build and behavior verification")
    build_meta = manifest["build"]
    if build_meta.get("command"):
        paragraph(doc, "Verification command: " + english_text(build_meta["command"]))
    fallback = "Passed" if build_meta.get("passed") else "Verification has not completed"
    paragraph(doc, "Result: " + english_text(build_meta.get("summary", fallback)))
    if build_meta.get("tests") is not None:
        paragraph(doc, "Tests executed: " + str(build_meta["tests"]))
    checks = manifest.get("checks", [])
    if checks:
        table(doc, ["Check", "Observed result"],
              [[english_text(x.get("name", x.get("case", ""))),
                english_text(x.get("result", x.get("details", "")))] for x in checks], [5.2, 11.8])
    else:
        paragraph(doc, "The screenshot descriptions record the browser observations, including field errors, "
                       "retained input and the redirect after a valid submission.")
    if manifest.get("prg"):
        paragraph(doc, "Post/Redirect/Get verification: " + english_text(manifest["prg"]))

    heading(doc, "Responses for a missing book")
    paragraph(doc, "Both requests use the missing identifier 999. BookService throws BookNotFoundException; "
                   "the controller type and the scope of its advice determine the response format.")
    table(doc, ["Request", "Status and Content Type", "Handler and response"],
          [[english_text(x["request"]), f"{x['status']}\n{x['content_type']}",
            english_text(x.get("handler", "")) + "\n" + english_text(x.get("result", ""))]
           for x in responses], [4.4, 4.8, 7.8])

    heading(doc, "Application screenshots and observed results")
    paragraph(doc, "These screenshots were captured from the running application. The four required cases show "
                   "an empty title, the year 1200, a missing genre and a successful submission. "
                   "The English and Russian forms each include a validation message. "
                   "Further captures show @NoDigits, simultaneous field errors, 404 responses and Variant 11.")
    for i, record in enumerate(screenshots, 1):
        evidence_image(doc, record, manifest_path.parent, i)

    doc.add_page_break()
    heading(doc, "Source code screenshots")
    paragraph(doc, "The following listings are screenshots of the implemented source code. "
                   "Each listing identifies its repository path and original line ranges. "
                   "Long files are split into consecutive captures, including both complete resource bundles.")

    heading(doc, "Form object and data binding")
    paragraph(doc, "BookForm separates incoming fields from the stored Book record. "
                   "Integer distinguishes an empty year from a numeric value. @NotNull checks that the year "
                   "is present, while @Min and @Max enforce the allowed range from 1450 to 2100.")
    add_listing("The BookForm class", "BookForm.java", 1)
    add_listing("The Genre enumeration", "Genre.java", 2)

    heading(doc, "The form template")
    paragraph(doc, "th:object selects form as the object for *{...} expressions. "
                   "th:field generates the name, id and current field value, including when invalid input is redisplayed. "
                   "Messages appear beside the fields and in a summary; CSS highlights invalid fields.")
    add_listing("The complete book creation template", "form.html", 3)

    heading(doc, "Handling invalid and valid submissions")
    paragraph(doc, "@Valid triggers validation and Spring records violations in BindingResult. "
                   "When errors exist, the controller returns books/form and restores genres. "
                   "A valid POST finishes with redirect:/books. The localized flash message is available "
                   "only to the following request.")
    add_listing("The POST handler", "BookPageController.java", 4, "post")

    heading(doc, "The custom author constraint")
    paragraph(doc, "@NoDigits checks that an author name contains no digits. The annotation declares "
                   "the required message, groups and payload members and selects NoDigitsValidator. "
                   "Bloch 2018 is rejected with the message resolved from book.author.nodigits.")
    add_listing("The NoDigits annotation", "NoDigits.java", 5)
    add_listing("The NoDigitsValidator implementation", "NoDigitsValidator.java", 6)

    heading(doc, "Locale selection and centralized exception handling")
    paragraph(doc, "SessionLocaleResolver stores the selected language in the HTTP session. "
                   "LocaleChangeInterceptor processes ?lang=en and ?lang=ru. "
                   "PageExceptionHandler is scoped to the book page controller, while ApiExceptionHandler "
                   "handles @RestController classes. The same service exception therefore produces "
                   "the response expected by each client.")
    add_listing("Locale configuration", "LocaleConfiguration.java", 7)
    add_listing("HTML page advice and Variant 11", "PageExceptionHandler.java", 8)
    add_listing("API advice", "ApiExceptionHandler.java", 9)

    heading(doc, "Individual Variant 11")
    paragraph(doc, "Variant 11 requires @ControllerAdvice with @ModelAttribute, a shared model attribute "
                   "and an exception handler. In Listing 8, libraryName returns the localized library name. "
                   "The bookNotFound handler in the same class returns error/not-found with status 404.")
    paragraph(doc, "The shared heading is visible on the form, catalogue, detail page and HTML error page. "
                   "The normal GET handlers do not each add this attribute manually. "
                   "Selecting Russian also changes libraryName, as shown in Figures 4, 9 and 11.")

    heading(doc, "Resource bundles")
    paragraph(doc, "A constraint message in braces is interpreted as a resource key. MessageSource resolves "
                   "the text for the selected locale, BindingResult retains the error, and th:errors displays it. "
                   "messages.properties supplies the default bundle and fallback values. "
                   "Both files use UTF-8; the Russian values are retained to demonstrate the required second language.")
    add_listing("English labels and messages", "messages.properties", 10)
    add_listing("Russian labels and messages", "messages_ru.properties", 11)

    heading(doc, "Comparing the form with the stored model")
    paragraph(doc, "BookForm has setters for partial binding and constraints for incoming values. "
                   "Book represents a saved book and contains an identifier assigned by the service. "
                   "The Book record has no setters; its constructor receives values after successful form validation.")
    add_listing("The Book record", "Book.java", 12)
    paragraph(doc, "GET /books/new puts an empty BookForm into the model as form and adds Genre.values(). "
                   "This allows th:object to resolve during the first request and supplies the genre options.")
    add_listing("The GET handler for a new book", "BookPageController.java", 13, "get_new")

    heading(doc, "Conclusions")
    paragraph(doc, "Declarative validation places field rules in BookForm instead of repeating manual checks "
                   "in the controller. BindingResult allows all violations to be shown while retaining the submitted values. "
                   "Resource bundles separate text from code and keep labels and validation messages in the same language. "
                   "Centralized advice makes HTML and API errors consistent, while Post/Redirect/Get prevents "
                   "a result-page refresh from repeating the POST.")

    heading(doc, "Short answers for the demonstration")
    qa = [
        ("Why use a separate form object",
         "It defines the allowed input fields and validation rules. The stored model can have different fields and constraints."),
        ("Why place BindingResult immediately after the form",
         "Spring associates that result with the preceding bound object. Moving it can prevent normal handling of invalid input in this method."),
        ("What changes without @Valid",
         "The handler no longer invokes Bean Validation constraints on the object, although binding can still report type conversion errors."),
        ("How the required-value constraints differ",
         "@NotNull rejects null. @NotEmpty also rejects empty strings or collections. @NotBlank rejects empty or whitespace-only strings."),
        ("Which types of errors are used",
         "The form constraints produce field errors. A year conversion error also belongs to a field. A class-level constraint could create a global error."),
        ("Why NoDigitsValidator accepts null",
         "@NotBlank checks required input; the custom validator has the separate responsibility of rejecting digits."),
        ("How an exception handler is selected",
         "Spring considers the advice scope and exception type. Local controller handlers are considered before advice; advice order and match specificity affect other choices."),
        ("Why display user input with th:text",
         "th:text escapes HTML. th:utext outputs unescaped HTML, so using it for untrusted user input can introduce XSS."),
    ]
    for question, answer in qa:
        p = paragraph(doc, "")
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_after = Pt(4)
        p.add_run(question + ". ").bold = True
        p.add_run(answer)
        for run in p.runs:
            run.font.size = Pt(10.5)

    # English narration is mandatory; Russian UI/source values belong only in images.
    narrative = "\n".join(p.text for p in doc.paragraphs)
    narrative += "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    if re.search(r"[\u0400-\u04ff]", narrative):
        raise ValueError("A non-English narrative passage remains in the report.")
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(json.dumps({"output": str(output), "application_screenshots": len(screenshots),
                      "code_screenshots": sum(len(x) for x in code_manifest.values()),
                      "listings": 13, "responses": len(responses)}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--code-manifest", type=Path)
    parser.add_argument("--output", type=Path, default=REPORT_DIR / "Laboratory_work_6_report.docx")
    args = parser.parse_args()
    build(args.manifest.resolve(), args.output.resolve(),
          args.code_manifest.resolve() if args.code_manifest else None)
